from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


@dataclass
class LoadedDocument:
    filename: str
    text: str
    source_type: str
    metadata: dict = field(default_factory=dict)


def _clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def load_txt(path: Path) -> List[LoadedDocument]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [LoadedDocument(path.name, _clean_text(text), "txt")]


def load_pdf(path: Path) -> List[LoadedDocument]:
    reader = PdfReader(str(path))
    documents: List[LoadedDocument] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = _clean_text(page.extract_text() or "")
        if text:
            documents.append(
                LoadedDocument(
                    filename=path.name,
                    text=text,
                    source_type="pdf",
                    metadata={"page": page_number},
                )
            )
    return documents


def load_docx(path: Path) -> List[LoadedDocument]:
    doc = DocxDocument(str(path))
    paragraphs = [_clean_text(p.text) for p in doc.paragraphs if _clean_text(p.text)]
    text = "\n".join(paragraphs)
    return [LoadedDocument(path.name, text, "docx")]


def load_document(path: Path) -> List[LoadedDocument]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return load_txt(path)
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def load_documents(folder: Path) -> List[LoadedDocument]:
    folder.mkdir(parents=True, exist_ok=True)
    results: List[LoadedDocument] = []
    for path in sorted(folder.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                results.extend(load_document(path))
            except Exception as exc:
                print(f"Skipping {path.name}: {exc}")
    return [doc for doc in results if doc.text.strip()]
